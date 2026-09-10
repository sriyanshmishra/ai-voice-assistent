from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    HRFlowable,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "submission"
OUTPUT.mkdir(exist_ok=True)

NAVY = RGBColor(12, 31, 52)
TEAL = RGBColor(0, 153, 153)
MINT = RGBColor(224, 247, 241)
ORANGE = RGBColor(240, 126, 54)
INK = RGBColor(31, 43, 55)
LIGHT = RGBColor(244, 248, 249)


def add_pdf_page_number(canvas, document):
    canvas.saveState()
    canvas.setFillColor(colors.HexColor("#607080"))
    canvas.setFont("Helvetica", 8)
    canvas.drawRightString(195 * mm, 12 * mm, f"Offline Voice & Text AI Assistant  |  {document.page}")
    canvas.restoreState()


def build_pdf():
    path = OUTPUT / "Offline_Voice_AI_Project_Description.pdf"
    doc = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        rightMargin=18 * mm,
        leftMargin=18 * mm,
        topMargin=17 * mm,
        bottomMargin=19 * mm,
        title="Offline Voice & Text AI Assistant - Project Description",
        author="Sriyansh Mishra",
    )
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(
        name="CoverTitle", parent=styles["Title"], fontName="Helvetica-Bold",
        fontSize=26, leading=31, textColor=colors.HexColor("#0C1F34"), alignment=TA_CENTER,
        spaceAfter=10,
    ))
    styles.add(ParagraphStyle(
        name="Subtitle", parent=styles["Normal"], fontName="Helvetica",
        fontSize=12, leading=18, textColor=colors.HexColor("#3B5164"), alignment=TA_CENTER,
    ))
    styles.add(ParagraphStyle(
        name="Section", parent=styles["Heading2"], fontName="Helvetica-Bold",
        fontSize=17, leading=21, textColor=colors.HexColor("#0C1F34"), spaceBefore=8, spaceAfter=8,
    ))
    styles.add(ParagraphStyle(
        name="BodyClean", parent=styles["BodyText"], fontName="Helvetica",
        fontSize=10.3, leading=15, textColor=colors.HexColor("#253746"), spaceAfter=7,
    ))
    styles.add(ParagraphStyle(
        name="Small", parent=styles["BodyText"], fontName="Helvetica",
        fontSize=8.7, leading=12, textColor=colors.HexColor("#536675"),
    ))
    styles.add(ParagraphStyle(
        name="CardTitle", parent=styles["Heading3"], fontName="Helvetica-Bold",
        fontSize=11, leading=14, textColor=colors.HexColor("#0C1F34"), spaceAfter=4,
    ))

    story = []
    story.append(Spacer(1, 24 * mm))
    story.append(Paragraph("OFFLINE VOICE & TEXT<br/>AI ASSISTANT", styles["CoverTitle"]))
    story.append(Paragraph("A privacy-first edge AI assistant for voice and text interaction", styles["Subtitle"]))
    story.append(Spacer(1, 14 * mm))
    story.append(HRFlowable(width="70%", thickness=2, color=colors.HexColor("#009999"), hAlign="CENTER"))
    story.append(Spacer(1, 12 * mm))
    cover = Table([
        [Paragraph("HACKATHON PROJECT BRIEF", styles["CardTitle"])],
        [Paragraph("Built with Python, Whisper speech-to-text, and a deployment path for Qualcomm AI Hub / Snapdragon NPU optimization.", styles["BodyClean"])],
    ], colWidths=[150 * mm])
    cover.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#E0F7F1")),
        ("BOX", (0, 0), (-1, -1), 0.7, colors.HexColor("#66B7AD")),
        ("LEFTPADDING", (0, 0), (-1, -1), 10), ("RIGHTPADDING", (0, 0), (-1, -1), 10),
        ("TOPPADDING", (0, 0), (-1, -1), 9), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    story.append(cover)
    story.append(Spacer(1, 35 * mm))
    story.append(Paragraph("Submitted by: Sriyansh Mishra", styles["Subtitle"]))
    story.append(Paragraph("Repository: github.com/sriyanshmishra/ai-voice-assistent", styles["Small"]))
    story.append(PageBreak())

    story.append(Paragraph("1. Executive Summary", styles["Section"]))
    story.append(Paragraph(
        "Offline Voice & Text AI Assistant is a lightweight edge-AI prototype that accepts either microphone audio or typed text. Voice input is transcribed locally with OpenAI Whisper Tiny, then passed to a simple intent-handler layer. The design keeps audio and text processing on-device after the model is cached, reducing dependency on cloud services and improving privacy for everyday interactions.",
        styles["BodyClean"],
    ))
    story.append(Paragraph("The project is intentionally structured for a Snapdragon deployment path: the current development fallback uses Hugging Face Whisper, while the same input contract can be connected to a Qualcomm AI Hub-exported model for NPU execution.", styles["BodyClean"]))

    story.append(Paragraph("2. Problem", styles["Section"]))
    problem_data = [
        [Paragraph("Cloud dependency", styles["CardTitle"]), Paragraph("Voice assistants commonly send recordings to remote servers, creating latency, connectivity, and privacy concerns.", styles["BodyClean"])],
        [Paragraph("Fragmented interaction", styles["CardTitle"]), Paragraph("Users need a quick way to switch between voice and text depending on environment, accessibility, or hardware availability.", styles["BodyClean"])],
        [Paragraph("Edge hardware opportunity", styles["CardTitle"]), Paragraph("Modern Snapdragon devices can run AI locally, but prototypes need a clear bridge from model experimentation to device-optimized inference.", styles["BodyClean"])],
    ]
    table = Table(problem_data, colWidths=[42 * mm, 116 * mm], hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#0C1F34")),
        ("TEXTCOLOR", (0, 0), (0, -1), colors.white),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#C9D7DE")),
        ("LEFTPADDING", (0, 0), (-1, -1), 8), ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(table)

    story.append(Paragraph("3. Solution", styles["Section"]))
    story.append(Paragraph("The assistant provides a compact, local-first interaction loop:", styles["BodyClean"]))
    solution = [
        ["1", "Choose voice or text input from the terminal interface."],
        ["2", "Record a short microphone sample or accept typed text."],
        ["3", "Transcribe voice locally with Whisper when voice mode is selected."],
        ["4", "Route the resulting text to a replaceable intent-handler function."],
        ["5", "Run from a cached model with WHISPER_LOCAL_ONLY=1 when offline operation is required."],
    ]
    solution_table = Table([[Paragraph(f"<b>{n}</b>", styles["BodyClean"]), Paragraph(text, styles["BodyClean"])] for n, text in solution], colWidths=[12 * mm, 146 * mm])
    solution_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F47E36")),
        ("TEXTCOLOR", (0, 0), (0, -1), colors.white),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LINEBELOW", (0, 0), (-1, -2), 0.3, colors.HexColor("#DCE5E8")),
        ("LEFTPADDING", (0, 0), (-1, -1), 8), ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 1),
    ]))
    story.append(solution_table)

    story.append(PageBreak())
    story.append(Paragraph("4. Technical Architecture", styles["Section"]))
    architecture = [
        [Paragraph("Input layer", styles["CardTitle"]), Paragraph("sounddevice microphone capture or terminal text input", styles["BodyClean"])],
        [Paragraph("Pre-processing", styles["CardTitle"]), Paragraph("16 kHz float32 audio passed to the Whisper processor", styles["BodyClean"])],
        [Paragraph("Inference", styles["CardTitle"]), Paragraph("Whisper Tiny local model with PyTorch inference mode; designed for replacement with a Qualcomm AI Hub optimized runtime", styles["BodyClean"])],
        [Paragraph("Application layer", styles["CardTitle"]), Paragraph("Transcribed text flows into process_text_input(), a clear extension point for commands and intents", styles["BodyClean"])],
        [Paragraph("Configuration", styles["CardTitle"]), Paragraph("Environment variables control model selection and local-only loading", styles["BodyClean"])],
    ]
    arch_table = Table(architecture, colWidths=[42 * mm, 116 * mm])
    arch_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#E0F7F1")),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#B7CFD1")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 8), ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(arch_table)
    story.append(Spacer(1, 5 * mm))
    story.append(Paragraph("5. What Works Today", styles["Section"]))
    for item in [
        "Voice recording from the default microphone with a configurable five-second duration.",
        "Whisper-based transcription and local text decoding.",
        "Direct text input for testing without microphone access.",
        "Lazy model loading, so text mode can be tested before downloading Whisper.",
        "Offline-only mode after the model has been downloaded and cached.",
    ]:
        story.append(Paragraph(f"- {item}", styles["BodyClean"]))

    story.append(Paragraph("6. Snapdragon / Qualcomm AI Hub Roadmap", styles["Section"]))
    roadmap = [
        ["Current prototype", "Whisper Tiny via Hugging Face + PyTorch development fallback"],
        ["Optimization", "Export and compile the compatible Whisper graph with Qualcomm AI Hub for a target Snapdragon device"],
        ["Runtime", "Replace the model invocation behind the existing transcription interface with the device runtime"],
        ["Productization", "Add intents, wake-word handling, streaming audio, UI, and on-device evaluation metrics"],
    ]
    roadmap_table = Table(roadmap, colWidths=[38 * mm, 120 * mm])
    roadmap_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#0C1F34")),
        ("TEXTCOLOR", (0, 0), (0, -1), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#C9D7DE")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("LEADING", (0, 0), (-1, -1), 12),
        ("LEFTPADDING", (0, 0), (-1, -1), 7), ("RIGHTPADDING", (0, 0), (-1, -1), 7),
        ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.append(roadmap_table)

    story.append(Paragraph("7. Expected Impact", styles["Section"]))
    story.append(Paragraph("The project demonstrates a practical path to private, low-latency AI interaction on edge hardware. It can support environments with weak connectivity, reduce the need to transmit sensitive audio, and provide a modular foundation for accessibility tools, field assistants, education, and device control.", styles["BodyClean"]))
    story.append(Paragraph("8. Repository", styles["Section"]))
    story.append(Paragraph("github.com/sriyanshmishra/ai-voice-assistent", styles["BodyClean"]))
    story.append(Spacer(1, 5 * mm))
    story.append(Paragraph("Note: The current repository is a working prototype. Qualcomm NPU acceleration is a planned integration path and should be validated on the target Snapdragon hardware before making performance claims.", styles["Small"]))
    doc.build(story, onFirstPage=add_pdf_page_number, onLaterPages=add_pdf_page_number)
    return path


def add_text(slide, text, left, top, width, height, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.05)
    frame.margin_right = Inches(0.05)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_footer(slide, number):
    add_text(slide, "OFFLINE VOICE & TEXT AI ASSISTANT", 0.55, 7.12, 5.5, 0.2, 7, RGBColor(90, 110, 125), True)
    add_text(slide, f"{number:02d}", 12.35, 7.12, 0.45, 0.2, 7, RGBColor(90, 110, 125), True, PP_ALIGN.RIGHT)


def add_header(slide, kicker, title, number):
    add_text(slide, kicker.upper(), 0.6, 0.38, 5.8, 0.22, 9, TEAL, True)
    add_text(slide, title, 0.6, 0.72, 11.4, 0.55, 25, NAVY, True)
    add_footer(slide, number)


def add_card(slide, left, top, width, height, title, body, accent=TEAL):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = LIGHT
    shape.line.color.rgb = RGBColor(214, 226, 230)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(0.08), Inches(height))
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent
    bar.line.fill.background()
    add_text(slide, title, left + 0.25, top + 0.16, width - 0.45, 0.34, 15, NAVY, True)
    add_text(slide, body, left + 0.25, top + 0.56, width - 0.45, height - 0.68, 11, INK)


def build_ppt():
    path = OUTPUT / "Offline_Voice_AI_Pitch_Deck.pptx"
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    slide = prs.slides.add_slide(blank)
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = NAVY
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.22), Inches(7.5))
    accent.fill.solid(); accent.fill.fore_color.rgb = TEAL; accent.line.fill.background()
    add_text(slide, "HACKATHON PITCH", 0.8, 1.05, 5, 0.3, 12, RGBColor(109, 226, 210), True)
    add_text(slide, "Offline Voice & Text\nAI Assistant", 0.8, 1.65, 8.4, 1.55, 34, RGBColor(255, 255, 255), True)
    add_text(slide, "Private, low-latency voice interaction designed for the edge", 0.85, 3.55, 7.8, 0.5, 17, RGBColor(214, 231, 236))
    orb = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(9.6), Inches(1.25), Inches(2.3), Inches(2.3))
    orb.fill.solid(); orb.fill.fore_color.rgb = TEAL; orb.line.fill.background()
    add_text(slide, "MIC\n→\nTEXT", 9.88, 1.77, 1.75, 1.25, 21, RGBColor(255, 255, 255), True, PP_ALIGN.CENTER)
    add_text(slide, "Sriyansh Mishra  |  ai-voice-assistent", 0.85, 6.55, 7, 0.25, 10, RGBColor(190, 211, 219))
    add_footer(slide, 1)

    slide = prs.slides.add_slide(blank)
    add_header(slide, "The challenge", "Voice AI should not require the cloud", 2)
    add_card(slide, 0.7, 1.65, 3.75, 2.15, "Privacy", "Sensitive audio should stay on the device whenever possible.", ORANGE)
    add_card(slide, 4.78, 1.65, 3.75, 2.15, "Connectivity", "Cloud-only assistants fail when networks are slow, unavailable, or restricted.", TEAL)
    add_card(slide, 8.86, 1.65, 3.75, 2.15, "Edge opportunity", "Snapdragon hardware can bring AI inference closer to the user.", NAVY)
    add_text(slide, "The opportunity: combine a simple interaction loop with a hardware-ready local inference path.", 1.05, 4.75, 11.1, 0.75, 22, NAVY, True, PP_ALIGN.CENTER)

    slide = prs.slides.add_slide(blank)
    add_header(slide, "The solution", "One assistant, two natural inputs", 3)
    add_card(slide, 0.8, 1.55, 5.45, 3.25, "1  Voice mode", "Record a short microphone sample, transcribe it locally with Whisper, and route the text to an intent handler.", TEAL)
    add_card(slide, 7.05, 1.55, 5.45, 3.25, "2  Text mode", "Type a request directly for quick testing, accessibility, or situations where a microphone is unavailable.", ORANGE)
    add_text(slide, "Both paths converge on the same text-processing layer.", 2.1, 5.55, 9.2, 0.45, 18, NAVY, True, PP_ALIGN.CENTER)

    slide = prs.slides.add_slide(blank)
    add_header(slide, "How it works", "A compact local-first pipeline", 4)
    steps = [("INPUT", "Mic or text", TEAL), ("PROCESS", "16 kHz audio", ORANGE), ("INFER", "Whisper", NAVY), ("ACT", "Intent handler", TEAL)]
    x = 0.7
    for index, (label, body, color) in enumerate(steps):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.05), Inches(2.45), Inches(1.55))
        shape.fill.solid(); shape.fill.fore_color.rgb = LIGHT; shape.line.color.rgb = color
        add_text(slide, label, x + 0.15, 2.27, 2.15, 0.28, 11, color, True, PP_ALIGN.CENTER)
        add_text(slide, body, x + 0.15, 2.7, 2.15, 0.35, 17, NAVY, True, PP_ALIGN.CENTER)
        if index < 3:
            add_text(slide, "->", x + 2.55, 2.52, 0.48, 0.45, 20, ORANGE, True, PP_ALIGN.CENTER)
        x += 3.15
    add_text(slide, "Offline mode is enabled after the model is downloaded and cached.", 1.15, 4.75, 11, 0.48, 19, NAVY, True, PP_ALIGN.CENTER)
    add_text(slide, "The model is loaded only when voice mode is selected, so text mode stays lightweight.", 1.65, 5.38, 10, 0.35, 12, INK, False, PP_ALIGN.CENTER)

    slide = prs.slides.add_slide(blank)
    add_header(slide, "Prototype status", "What is working today", 5)
    add_card(slide, 0.75, 1.45, 5.75, 1.25, "Voice capture", "Configurable five-second recording from the default microphone.", TEAL)
    add_card(slide, 6.8, 1.45, 5.75, 1.25, "Whisper transcription", "Local speech-to-text using Whisper Tiny and PyTorch.", ORANGE)
    add_card(slide, 0.75, 3.05, 5.75, 1.25, "Text testing", "Direct text mode works without downloading the Whisper model.", NAVY)
    add_card(slide, 6.8, 3.05, 5.75, 1.25, "Configurable", "Model path, sample rate, duration, and offline mode are configurable.", TEAL)
    add_text(slide, "Working prototype today  ->  Snapdragon-optimized runtime next", 1.05, 5.55, 11.2, 0.55, 21, NAVY, True, PP_ALIGN.CENTER)

    slide = prs.slides.add_slide(blank)
    add_header(slide, "Snapdragon path", "From development fallback to NPU inference", 6)
    add_card(slide, 0.7, 1.45, 3.72, 3.3, "01  Develop", "Validate the Whisper input/output contract with a CPU-friendly local model.", TEAL)
    add_card(slide, 4.8, 1.45, 3.72, 3.3, "02  Optimize", "Export and compile a compatible model with Qualcomm AI Hub for target Snapdragon hardware.", ORANGE)
    add_card(slide, 8.9, 1.45, 3.72, 3.3, "03  Deploy", "Swap the inference backend behind the same transcription interface and measure latency, memory, and accuracy.", NAVY)
    add_text(slide, "Honest status: the repository is NPU-ready by design; target-device acceleration still needs hardware validation.", 1.0, 5.45, 11.25, 0.65, 17, NAVY, True, PP_ALIGN.CENTER)

    slide = prs.slides.add_slide(blank)
    add_header(slide, "Why it matters", "Useful beyond the demo", 7)
    add_card(slide, 0.7, 1.55, 3.72, 2.6, "Privacy-first", "Keep everyday voice interactions local instead of sending recordings to a remote API.", TEAL)
    add_card(slide, 4.8, 1.55, 3.72, 2.6, "Resilient", "Continue working in low-connectivity environments and on restricted networks.", ORANGE)
    add_card(slide, 8.9, 1.55, 3.72, 2.6, "Extensible", "Add commands, device control, accessibility workflows, or domain-specific intents.", NAVY)
    add_text(slide, "Potential domains: accessibility | field work | education | device control | private personal assistants", 0.9, 5.1, 11.5, 0.5, 18, NAVY, True, PP_ALIGN.CENTER)

    slide = prs.slides.add_slide(blank)
    bg = slide.background.fill; bg.solid(); bg.fore_color.rgb = NAVY
    add_text(slide, "THE ASK", 0.8, 1.1, 4, 0.25, 11, RGBColor(109, 226, 210), True)
    add_text(slide, "Help us take private voice AI\nfrom prototype to Snapdragon edge hardware.", 0.8, 1.7, 10.8, 1.2, 31, RGBColor(255, 255, 255), True)
    add_text(slide, "Next milestone: Qualcomm AI Hub model export, target-device benchmarking, and a richer intent layer.", 0.85, 3.55, 10.8, 0.6, 17, RGBColor(214, 231, 236))
    add_text(slide, "Thank you", 0.85, 5.7, 4, 0.45, 24, RGBColor(109, 226, 210), True)
    add_text(slide, "github.com/sriyanshmishra/ai-voice-assistent", 0.85, 6.35, 8, 0.25, 11, RGBColor(190, 211, 219))
    add_footer(slide, 8)

    prs.save(path)
    return path


if __name__ == "__main__":
    print(build_pdf())
    print(build_ppt())